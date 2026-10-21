"""Excel model with live formulas. Raw real data in sheets; mapping, normalisation, scenario weights and ranks are formulas."""
import os, json, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.utils import get_column_letter as CL
import model as M
ROOT = M.ROOT; S = M.S.reset_index(); meta = json.load(open(os.path.join(ROOT, "outputs", "model_meta.json")))
gdp, loans, cos = M.gdp, M.loans.reset_index(), M.cos
ind = pd.read_csv(os.path.join(ROOT, "data", "knbs_economic_survey_indicators_2026.csv"))
F = "Arial"; HDR = PatternFill("solid", fgColor="12202B"); INP = Font(name=F, size=10, color="0000FF"); BASE = Font(name=F, size=10); B = Font(name=F, size=10, bold=True)
wb = Workbook()
def head(ws, r, vals, c0=1):
    for i, v in enumerate(vals):
        c = ws.cell(r, c0 + i, v); c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical="center")
def title(ws, t, sub=None):
    ws["A1"] = t; ws["A1"].font = Font(name=F, size=14, bold=True)
    if sub: ws["A2"] = sub; ws["A2"].font = Font(name=F, size=9, italic=True, color="555555")
def style(ws, rng, font=BASE, fmt=None):
    for row in ws[rng]:
        for c in row:
            c.font = font
            if fmt: c.number_format = fmt

# ---------------- README
ws = wb.active; ws.title = "README"
title(ws, "Kenya Industry Opportunity Radar: model", "Real data: KNBS Quarterly GDP Q1 2026, CBK Bank Supervision Annual Report 2025, CBK Credit Officer Survey (June 2026), KNBS Economic Survey 2026.")
lines = ["How to read: Scorecard ranks ten industry groups on five pillars. Change the scenario or weights on Inputs; everything recalculates.",
 "Blue cells are inputs; black cells are formulas. GDP and Credit_Map hold the editable mapping of activities and bank sectors to industry groups.",
 "Pillars: Scale (2025 real value added), Recent growth (Q1 2026 vs Q1 2025), Track record (2021-25 CAGR), Credit demand (net % of credit officers seeing higher demand, June 2026, loan-weighted), Credit health (inverse of NPL ratio, Dec 2025).",
 "Each pillar is min-max scaled to 0-100 across the ten groups. Groups with no matching bank sector (Public Sector & Development, Health) get a neutral 50 on the two credit pillars.",
 "Transport and Communication is one CBK bank sector; it is shared by Technology, Media & Telecom and Transport & Logistics.",
 "Industry groups are an analyst grouping for this exercise, not any firm's official taxonomy. Professional services, other services, FISIM and taxes on products are excluded.",
 "Checks sheet compares every headline number with the independent Python calculation (analysis/model.py)."]
for i, t in enumerate(lines): ws.cell(4 + i, 1, t).font = BASE
ws.column_dimensions["A"].width = 160

# ---------------- Inputs
wi = wb.create_sheet("Inputs"); title(wi, "Inputs and scenario weights", "Weights are % of the total score. Pick the scenario used on the Scorecard.")
wi["A3"] = "Selected scenario"; wi["B3"] = "Balanced"; wi["A4"] = "Neutral score for missing credit pillar"; wi["B4"] = M.NEUTRAL
wi["A5"] = "Tier cut-off: Lead = rank at or below"; wi["B5"] = 3; wi["A6"] = "Tier cut-off: Grow = rank at or below"; wi["B6"] = 7
for c in ("B3", "B4", "B5", "B6"): wi[c].font = INP
for r in range(3, 7): wi.cell(r, 1).font = B
head(wi, 8, ["Scenario", "Scale", "Recent growth", "Track record", "Credit demand", "Credit health", "Total"])
for i, (k, w) in enumerate(meta["scenarios"].items()):
    r = 9 + i; wi.cell(r, 1, k).font = BASE
    for j, x in enumerate(w): c = wi.cell(r, 2 + j, x); c.font = INP
    wi.cell(r, 7, f"=SUM(B{r}:F{r})").font = BASE
dv = DataValidation(type="list", formula1="=$A$9:$A$13"); wi.add_data_validation(dv); dv.add("B3")
wi.column_dimensions["A"].width = 40
for c in "BCDEFG": wi.column_dimensions[c].width = 15

# ---------------- GDP (real, constant 2016 prices, KSh m)
wg = wb.create_sheet("GDP"); title(wg, "KNBS real GDP by activity (constant 2016 prices, KSh million)", "Source: KNBS Quarterly Gross Domestic Product Report, First Quarter 2026, Table 2. Column B is the editable group mapping.")
head(wg, 4, ["Activity", "Industry group", "FY2021", "FY2024", "FY2025", "Q1 2025", "Q1 2026", "Growth 2025 %", "Growth Q1 2026 %"])
acts = [a for a in gdp.activity.unique() if a not in ("All Industries at Basic Prices", "Taxes on Products", "GDP at Market Prices", "GDP, Seasonally Adjusted")]
amap = {a: g for g, (aa, _) in M.GROUPS.items() for a in aa}
def gv(a, y, p): return float(gdp[(gdp.activity == a) & (gdp.year == y) & (gdp.period == p)].value.iloc[0])
for i, a in enumerate(acts):
    r = 5 + i; wg.cell(r, 1, a); wg.cell(r, 2, amap.get(a, "Excluded")).font = INP
    for j, (y, p) in enumerate([(2021, "FY"), (2024, "FY"), (2025, "FY"), (2025, "Q1"), (2026, "Q1")]): wg.cell(r, 3 + j, gv(a, y, p)).font = INP
    wg.cell(r, 8, f"=E{r}/D{r}*100-100"); wg.cell(r, 9, f"=G{r}/F{r}*100-100")
GL = 4 + len(acts); rt = GL + 1
wg.cell(rt, 1, "All industries at basic prices (published)").font = B
for j, (y, p) in enumerate([(2021, "FY"), (2024, "FY"), (2025, "FY"), (2025, "Q1"), (2026, "Q1")]): wg.cell(rt, 3 + j, gv("All Industries at Basic Prices", y, p)).font = INP
wg.cell(rt + 1, 1, "Sum of activities listed (excludes FISIM adjustment)").font = BASE
for j in range(5): wg.cell(rt + 1, 3 + j, f"=SUM({CL(3+j)}5:{CL(3+j)}{GL})")
wg.cell(rt + 2, 1, "Difference vs published (FISIM)").font = BASE
for j in range(5): wg.cell(rt + 2, 3 + j, f"={CL(3+j)}{rt}-{CL(3+j)}{rt+1}")
style(wg, f"C5:G{rt+2}", BASE, "#,##0"); style(wg, f"H5:I{GL}", BASE, "0.0")
for c, w in zip("ABCDEFGHI", [42, 30, 14, 14, 14, 14, 14, 14, 16]): wg.column_dimensions[c].width = w
wg.freeze_panes = "C5"

# ---------------- Credit raw
wc = wb.create_sheet("Credit_Raw"); title(wc, "CBK bank lending, NPLs and credit officer survey by sector", "Sources: CBK Bank Supervision Annual Report 2025, Table 18 (Dec 2025); CBK Credit Officer Survey, quarter ended June 2026 (% of respondents).")
head(wc, 4, ["Bank sector", "Loan accounts", "Gross loans KSh m", "Gross NPLs KSh m", "NPL ratio %", "Demand: increased %", "Demand: unchanged %", "Demand: decreased %", "Standards: tightened %", "Standards: eased %", "NPL outlook: increase %", "NPL outlook: decrease %"])
def cv(m, bs, col): return int(cos[(cos.measure == m) & (cos.period == "2026-06") & (cos.bank_sector == bs)][col].iloc[0])
for i, row in loans.reset_index(drop=True).iterrows():
    r = 5 + i; bs = row.bank_sector
    vals = [bs, row.loan_accounts, row.gross_loans_ksh_m, row.gross_npl_ksh_m, f"=D{r}/C{r}*100", cv("demand", bs, "up_pct"), cv("demand", bs, "same_pct"), cv("demand", bs, "down_pct"), cv("standards", bs, "up_pct"), cv("standards", bs, "down_pct"), cv("npl_expect", bs, "up_pct"), cv("npl_expect", bs, "down_pct")]
    for j, v in enumerate(vals): wc.cell(r, 1 + j, v).font = BASE if j in (0, 4) else INP
CE = 4 + len(loans); wc.cell(CE + 1, 1, "Total").font = B
for j in (2, 3, 4): wc.cell(CE + 1, j, f"=SUM({CL(j)}5:{CL(j)}{CE})").font = B
wc.cell(CE + 1, 5, f"=D{CE+1}/C{CE+1}*100").font = B
style(wc, f"B5:D{CE+1}", BASE, "#,##0"); style(wc, f"E5:E{CE+1}", BASE, "0.0")
wc.column_dimensions["A"].width = 32
for j in range(2, 13): wc.column_dimensions[CL(j)].width = 15

# ---------------- Credit map (group -> bank sector)
wm = wb.create_sheet("Credit_Map"); title(wm, "Mapping of industry groups to CBK bank sectors", "Editable (blue). Transport and Communication is shared by two groups.")
head(wm, 4, ["Industry group", "Bank sector", "Gross loans", "Gross NPLs", "Demand up %", "Demand down %"])
mp = [(g, b) for g, (_, bb) in M.GROUPS.items() for b in bb]
for i, (g, b) in enumerate(mp):
    r = 5 + i; wm.cell(r, 1, g).font = INP; wm.cell(r, 2, b).font = INP
    for j, col in enumerate("CDFH"):
        wm.cell(r, 3 + j, f"=INDEX(Credit_Raw!${col}$5:${col}${CE},MATCH($B{r},Credit_Raw!$A$5:$A${CE},0))").font = BASE
ME = 4 + len(mp); style(wm, f"C5:F{ME}", BASE, "#,##0")
wm.column_dimensions["A"].width = 32; wm.column_dimensions["B"].width = 32
for c in "CDEF": wm.column_dimensions[c].width = 14

# ---------------- Scorecard
sc = wb.create_sheet("Scorecard", 1); title(sc, "Industry Opportunity Radar: scorecard", "Real KNBS and CBK data. Scenario selected on Inputs. Ranks 1 = highest priority.")
sc["A3"] = "Scenario:"; sc["B3"] = "=Inputs!B3"; sc["A3"].font = B; sc["B3"].font = Font(name=F, size=10, bold=True, color="C0392B")
H = ["Industry group", "Value added 2025 (KSh m)", "Share of GDP %", "Growth 2025 %", "CAGR 2021-25 %", "Growth Q1 2026 %", "Gross loans (KSh m)", "NPL ratio %", "Credit demand net %", "Scale (0-100)", "Recent growth", "Track record", "Credit demand", "Credit health", "Score", "Rank", "Tier", "Evidence"]
head(sc, 5, H); sc.row_dimensions[5].height = 42
G0, G1 = 6, 5 + len(M.GROUPS)
for i, g in enumerate(M.GROUPS):
    r = G0 + i; sc.cell(r, 1, g).font = B
    ga = f"GDP!$B$5:$B${GL}"
    sc.cell(r, 2, f"=SUMIFS(GDP!$E$5:$E${GL},{ga},$A{r})")
    sc.cell(r, 3, f"=B{r}/GDP!$E${rt}*100")
    sc.cell(r, 4, f"=B{r}/SUMIFS(GDP!$D$5:$D${GL},{ga},$A{r})*100-100")
    sc.cell(r, 5, f"=((B{r}/SUMIFS(GDP!$C$5:$C${GL},{ga},$A{r}))^(1/4)-1)*100")
    sc.cell(r, 6, f"=SUMIFS(GDP!$G$5:$G${GL},{ga},$A{r})/SUMIFS(GDP!$F$5:$F${GL},{ga},$A{r})*100-100")
    sc.cell(r, 7, f'=IF(SUMIFS(Credit_Map!$C$5:$C${ME},Credit_Map!$A$5:$A${ME},$A{r})=0,"",SUMIFS(Credit_Map!$C$5:$C${ME},Credit_Map!$A$5:$A${ME},$A{r}))')
    sc.cell(r, 8, f'=IF(G{r}="","",SUMIFS(Credit_Map!$D$5:$D${ME},Credit_Map!$A$5:$A${ME},$A{r})/G{r}*100)')
    sc.cell(r, 9, f'=IF(G{r}="","",SUMPRODUCT((Credit_Map!$A$5:$A${ME}=$A{r})*Credit_Map!$C$5:$C${ME}*(Credit_Map!$E$5:$E${ME}-Credit_Map!$F$5:$F${ME}))/G{r})')
    def mmx(col, inv=False):
        rg = f"${col}${G0}:${col}${G1}"
        z = f'(({col}{r}-MIN({rg}))/(MAX({rg})-MIN({rg}))*100)'
        return f'=IF({col}{r}="","",{"100-" if inv else ""}{z})'
    sc.cell(r, 10, mmx("B")); sc.cell(r, 11, mmx("F")); sc.cell(r, 12, mmx("E")); sc.cell(r, 13, mmx("I")); sc.cell(r, 14, mmx("H", True))
    f = lambda col: f'IF(ISNUMBER({col}{r}),{col}{r},Inputs!$B$4)'
    w = "INDEX(Inputs!$B$9:$F$13,MATCH(Inputs!$B$3,Inputs!$A$9:$A$13,0),{})"
    terms = "+".join(f"{f(c)}*{w.format(k+1)}" for k, c in enumerate("JKLMN"))
    sc.cell(r, 15, f"=({terms})/INDEX(Inputs!$G$9:$G$13,MATCH(Inputs!$B$3,Inputs!$A$9:$A$13,0))")
    sc.cell(r, 16, f"=RANK(O{r},$O${G0}:$O${G1})")
    sc.cell(r, 17, f'=IF(P{r}<=Inputs!$B$5,"Lead",IF(P{r}<=Inputs!$B$6,"Grow","Selective"))')
    sc.cell(r, 18, f'=IF(G{r}="","Growth only (credit neutral)","Growth + credit")')
style(sc, f"A{G0}:R{G1}", BASE); 
for r in range(G0, G1 + 1): sc.cell(r, 1).font = B
style(sc, f"B{G0}:B{G1}", BASE, "#,##0"); style(sc, f"C{G0}:F{G1}", BASE, "0.0"); style(sc, f"G{G0}:G{G1}", BASE, "#,##0"); style(sc, f"H{G0}:I{G1}", BASE, "0.0"); style(sc, f"J{G0}:O{G1}", BASE, "0.0"); style(sc, f"P{G0}:Q{G1}", BASE, "0")
sc.conditional_formatting.add(f"O{G0}:O{G1}", ColorScaleRule(start_type="min", start_color="F8D7D0", end_type="max", end_color="BFE3D0"))
sc.conditional_formatting.add(f"Q{G0}:Q{G1}", CellIsRule(operator="equal", formula=['"Lead"'], fill=PatternFill("solid", bgColor="BFE3D0")))
for c, w in zip([CL(i) for i in range(1, 19)], [30, 16, 11, 11, 12, 12, 15, 10, 12, 10, 10, 10, 10, 10, 9, 8, 11, 26]): sc.column_dimensions[c].width = w
sc.freeze_panes = "B6"
ch = BarChart(); ch.type = "bar"; ch.title = "Opportunity score by industry group (selected scenario)"; ch.height = 9; ch.width = 22
ch.add_data(Reference(sc, min_col=15, min_row=5, max_row=G1), titles_from_data=True); ch.set_categories(Reference(sc, min_col=1, min_row=G0, max_row=G1)); ch.legend = None; ch.y_axis.scaling.min = 0
sc.add_chart(ch, f"A{G1+3}")
n = G1 + 22
sc.cell(n, 1, "Reading the table: growth and size say where the economy is moving; credit demand and NPLs say where banks see client activity and risk. Public Sector and Health have no matching bank sector, so their credit pillars are neutral.").font = Font(name=F, size=9, italic=True)

# ---------------- Scenarios
wsn = wb.create_sheet("Scenarios", 2); title(wsn, "Scores and ranks under five weighting scenarios", "A robust priority stays near the top whatever the weights. Rank range shows sensitivity.")
names = list(meta["scenarios"].keys())
head(wsn, 4, ["Industry group"] + [f"Score: {k}" for k in names] + [f"Rank: {k}" for k in names] + ["Best rank", "Worst rank", "Stability"])
for i, g in enumerate(M.GROUPS):
    r = 5 + i; sr = G0 + i; wsn.cell(r, 1, f"=Scorecard!A{sr}").font = B
    for k, nm in enumerate(names):
        wr = 9 + k
        f = lambda col: f"IF(ISNUMBER(Scorecard!{col}{sr}),Scorecard!{col}{sr},Inputs!$B$4)"
        terms = "+".join(f"{f(c)}*Inputs!{CL(2+q)}${wr}" for q, c in enumerate("JKLMN"))
        wsn.cell(r, 2 + k, f"=({terms})/Inputs!$G${wr}")
        wsn.cell(r, 7 + k, f"=RANK({CL(2+k)}{r},{CL(2+k)}$5:{CL(2+k)}${4+len(M.GROUPS)})")
    wsn.cell(r, 12, f"=MIN(G{r}:K{r})"); wsn.cell(r, 13, f"=MAX(G{r}:K{r})"); wsn.cell(r, 14, f'=IF(M{r}-L{r}<=3,"Stable","Sensitive")')
style(wsn, f"B5:F{4+len(M.GROUPS)}", BASE, "0.0"); style(wsn, f"G5:N{4+len(M.GROUPS)}", BASE, "0")
wsn.column_dimensions["A"].width = 30
for j in range(2, 15): wsn.column_dimensions[CL(j)].width = 13

# ---------------- Signals
wsg = wb.create_sheet("Signals"); title(wsg, "Real-economy signals, 2021-2025", "Source: KNBS Economic Survey 2026, Annex 1 Table 0 (2025 provisional).")
head(wsg, 4, ["Indicator", "Unit", "2021", "2022", "2023", "2024", "2025", "Growth 2025 %", "CAGR 2021-25 %"])
for i, row in ind.reset_index(drop=True).iterrows():
    r = 5 + i; wsg.cell(r, 1, row.indicator); wsg.cell(r, 2, row.unit)
    for j, y in enumerate(["2021", "2022", "2023", "2024", "2025"]): wsg.cell(r, 3 + j, float(row[y])).font = INP
    wsg.cell(r, 8, f"=G{r}/F{r}*100-100"); wsg.cell(r, 9, f"=((G{r}/C{r})^(1/4)-1)*100")
style(wsg, f"C5:G{4+len(ind)}", BASE, "#,##0.0"); style(wsg, f"H5:I{4+len(ind)}", BASE, "0.0")
wsg.column_dimensions["A"].width = 40; wsg.column_dimensions["B"].width = 14
for j in range(3, 10): wsg.column_dimensions[CL(j)].width = 14

# ---------------- Python reference + Checks
wp = wb.create_sheet("Python_Ref"); title(wp, "Independent Python results (analysis/model.py)")
cols = ["group", "value_added_2025_ksh_m", "share_of_gdp_pct", "growth_2025_pct", "cagr_2021_25_pct", "growth_q1_2026_pct", "gross_loans_ksh_m", "npl_ratio_pct", "credit_demand_net_pct"] + [f"score_{k}" for k in names] + [f"rank_{k}" for k in names]
head(wp, 3, cols)
Sd = S.set_index("group")
for i, g in enumerate(M.GROUPS):
    r = 4 + i; wp.cell(r, 1, g)
    for j, c in enumerate(cols[1:]):
        v = Sd.loc[g, c]; wp.cell(r, 2 + j, None if pd.isna(v) else float(v))
wk = wb.create_sheet("Checks"); title(wk, "Reconciliation: workbook formulas vs Python", "Absolute tolerance 0.001 (levels 0.5).")
head(wk, 3, ["Check", "Workbook", "Python", "Abs diff", "Result"])
rowk = 4
def add(label, wbf, pyf, tol=0.001):
    global rowk
    wk.cell(rowk, 1, label); wk.cell(rowk, 2, wbf); wk.cell(rowk, 3, pyf); wk.cell(rowk, 4, f"=ABS(B{rowk}-C{rowk})"); wk.cell(rowk, 5, f'=IF(D{rowk}<={tol},"PASS","FAIL")'); rowk += 1
pc = {c: CL(2 + j) for j, c in enumerate(cols[1:])}
for i, g in enumerate(M.GROUPS):
    sr, pr = G0 + i, 4 + i
    add(f"{g}: value added", f"=Scorecard!B{sr}", f"=Python_Ref!{pc['value_added_2025_ksh_m']}{pr}", 0.5)
    for lab, scol, pcol in [("growth 2025", "D", "growth_2025_pct"), ("CAGR", "E", "cagr_2021_25_pct"), ("Q1 2026 growth", "F", "growth_q1_2026_pct")]:
        add(f"{g}: {lab}", f"=Scorecard!{scol}{sr}", f"=Python_Ref!{pc[pcol]}{pr}")
    if not pd.isna(Sd.loc[g, "npl_ratio_pct"]):
        add(f"{g}: NPL ratio", f"=Scorecard!H{sr}", f"=Python_Ref!{pc['npl_ratio_pct']}{pr}"); add(f"{g}: credit demand net", f"=Scorecard!I{sr}", f"=Python_Ref!{pc['credit_demand_net_pct']}{pr}")
    for k, nm in enumerate(names):
        add(f"{g}: score {nm}", f"=Scenarios!{CL(2+k)}{5+i}", f"=Python_Ref!{pc['score_'+nm]}{pr}", 0.01)
        add(f"{g}: rank {nm}", f"=Scenarios!{CL(7+k)}{5+i}", f"=Python_Ref!{pc['rank_'+nm]}{pr}", 0)
add("Sum of 11 bank-sector loans = CBK published total (4,347,164.1)", f"=Credit_Raw!C{CE+1}", 4347164.1, 0.5)
add("Sum of NPLs = CBK published total (696,931.0)", f"=Credit_Raw!D{CE+1}", 696931.0, 0.5)
add("Aggregate NPL ratio = 16.0% (CBK Table 20)", f"=ROUND(Credit_Raw!E{CE+1},1)", 16.0, 0.001)
add("FY2025 all-industries growth vs KNBS Table 3 (4.8%)", f"=ROUND((GDP!E{rt}/GDP!D{rt}-1)*100,1)", 4.8, 0.05)
add("Q1 2026 all-industry growth vs KNBS 5.4%", f"=ROUND((GDP!G{rt}/GDP!F{rt}-1)*100,1)", 5.4, 0.05)
wk.cell(rowk + 1, 1, "OVERALL").font = B; wk.cell(rowk + 1, 5, f'=IF(COUNTIF(E4:E{rowk-1},"FAIL")=0,"ALL CHECKS PASS","CHECK FAILURES")').font = B
style(wk, f"A4:E{rowk-1}", BASE); style(wk, f"B4:D{rowk-1}", BASE, "#,##0.000")
wk.column_dimensions["A"].width = 62
for c in "BCDE": wk.column_dimensions[c].width = 18
out = os.path.join(ROOT, "model", "opportunity_radar_model.xlsx"); os.makedirs(os.path.dirname(out), exist_ok=True); wb.save(out); print("saved", out, "checks", rowk - 4)
