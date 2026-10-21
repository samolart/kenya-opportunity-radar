"""Industry briefs, thought-leadership article, social post copy and content calendar, all generated from the real-data scorecard."""
import os, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import model as M
ROOT = M.ROOT; S = M.S.copy()
I = pd.read_csv(os.path.join(ROOT, "data", "knbs_economic_survey_indicators_2026.csv")).set_index("indicator")
gi = lambda k: I.loc[k, "growth_2025_pct"]
os.makedirs(os.path.join(ROOT, "briefs"), exist_ok=True); os.makedirs(os.path.join(ROOT, "comms"), exist_ok=True)
SRC = "Sources: KNBS Quarterly GDP Q1 2026 (Table 2); CBK Bank Supervision Annual Report 2025 (Table 18); CBK Credit Officer Survey, June 2026; KNBS Economic Survey 2026 (Annex 1)."
def f(x, d=1): return "n/a" if pd.isna(x) else f"{x:.{d}f}"
ctx = {
 "Financial Services": (f"Lending is growing (total domestic credit +{gi('Total domestic credit'):.1f}% in 2025) and the sector has the cleanest book among bank sectors. Mobile money transfers fell {abs(gi('Total mobile money transfer')):.1f}% in 2025, a sign that payments are changing shape.",
   ["How are you growing lending while keeping bad loans low?", "Where do you see digital payments moving, given mobile money value fell in 2025?"], "Lead with risk, regulation and digital operating-model topics."),
 "Consumer & Retail": (f"Trade shows the strongest credit demand of any business sector (65% of banks report higher demand), yet 19.9% of trade loans are non-performing. Tourism and hotels lend at a 22.9% NPL ratio, and more banks see demand falling than rising (25% vs 19%).",
   ["Working capital is in demand: how are you managing supplier and stock financing?", "Tourist arrivals grew 6.2% to 2,550.6 thousand in 2025; how are you capturing that demand?"], "Pair growth content with credit-risk and cash-management insight."),
 "Public Sector & Development": (f"Public administration and education are 12.2% of the economy. Ordinary revenue and grants grew {gi('Ordinary Revenue and Grants'):.1f}% in 2025 against expenditure growth of {gi('Total National Government Expenditure'):.1f}%. External debt service was 4.9% of GDP, or 31.1% of export earnings.",
   ["How is your organisation planning around tighter fiscal space and higher debt service?", "Where could better public financial management and reporting free up capacity?"], "No bank-credit signal exists, so the rank rests on size and growth only."),
 "Energy & Resources": (f"Electricity consumption rose {gi('Electricity consumption'):.1f}% in 2025 to 11,785.5 GWh. Energy and water has one of the cleanest loan books (6.8% NPL). Mining grew 9.1% in Q1 2026 from a small base.",
   ["Demand for power is up almost 10%: what does that mean for your capacity and grid plans?", "How are you funding the next investment cycle?"], "Strong growth from a small base; good for focused, high-value conversations."),
 "Agribusiness": (f"Agriculture is the largest sector (18.3% of GDP) but grew only 3.1% in 2025 and 4.9% in Q1 2026. Tea marketed production fell {abs(gi('Tea-marketed production')):.1f}% while horticultural exports grew {gi('Fresh Horticultural Produce exports'):.1f}% and milk sold centrally {gi('Milk sold centrally'):.1f}%. Bank NPLs are 19.3%.",
   ["Horticulture exports are growing while tea fell: how is your portfolio shifting?", "How are you protecting margins against weather and price swings?"], "Rank is sensitive to weights: high on scale, low on growth."),
 "Real Estate & Construction": (f"Cement consumption jumped {gi('Cement Consumption'):.1f}% and construction output rose {gi('Construction output'):.1f}% in 2025, but this group carries the weakest loan quality: 24.1% overall, 29.5% for building and construction alone. Credit officers see little change in demand.",
   ["Building activity is picking up: how are you managing payment delays and project finance?", "Which of your projects are most exposed to the high level of stressed loans in the sector?"], "Activity is up while credit risk is the highest; position around risk and project delivery."),
 "Technology, Media & Telecom": (f"Information and communication grew 7.8% a year over 2021 to 2025, the second-fastest track record. Mobile subscriptions reached 78.4 million (+{gi('Mobile Subscriptions'):.1f}%) but mobile money transfers fell {abs(gi('Total mobile money transfer')):.1f}%. Bank data covers transport and communication together (16.6% NPL), so credit signals are shared.",
   ["Subscriptions keep growing but mobile money value fell: where is the next revenue pool?", "How are you thinking about cybersecurity and data protection as services scale?"], "Strong long-run growth; credit read-across is imperfect."),
 "Health": (f"Health is 2.5% of GDP and grew 5.5% in 2025, with 4.3% growth in Q1 2026. Licensed doctors and dentists grew {gi('Licensed doctors and dentists'):.1f}% to 18,561. CBK reports no health bank sector, so the credit pillars are neutral.",
   ["How are you funding expansion in capacity and equipment?", "What would make universal coverage workable for your organisation?"], "Small and under-evidenced here; a selective, relationship-led play."),
 "Manufacturing": (f"Growth is the weakest of the ten groups (2.0% in 2025, 2.4% a year over 2021 to 2025). Output rose {gi('Manufacturing output'):.1f}% to KSh 3.8 trillion. Credit demand is solid (net +30) but 18.4% of loans are non-performing.",
   ["Demand for credit is rising: where are you investing, and what is holding back volumes?", "How are input costs and competitiveness affecting your margins?"], "Large and credit-hungry but slow-growing; select accounts, do not blanket-target."),
 "Transport & Logistics": (f"Rail freight grew {gi('Rail freight'):.1f}% to 8.4 million tonnes while air passengers fell {abs(gi('Air passengers handled')):.1f}%. Growth in Q1 2026 was 3.6%, the lowest of the ten groups.",
   ["Rail freight is up 11%: what does that mean for your logistics costs?", "How are you adjusting to flat air traffic?"], "Lowest score; keep to selective, account-driven outreach."),
}
out = ["# Industry briefs: Kenya Opportunity Radar", "", "*One short brief per industry group, generated from real KNBS and CBK data. Balanced weights unless stated. " + SRC + "*", ""]
for g, r in S.iterrows():
    why, qs, so = ctx[g]
    out += [f"## {int(r.rank_Balanced)}. {g} ({r.tier}; score {r.score_Balanced:.0f}/100; {r.robust.lower()} across weights)", "",
            f"- Size: {r.share_of_gdp_pct:.1f}% of GDP. Growth: {f(r.growth_2025_pct)}% in 2025, {f(r.growth_q1_2026_pct)}% in Q1 2026, {f(r.cagr_2021_25_pct)}% a year over 2021 to 2025.",
            f"- Bank credit: NPL ratio {f(r.npl_ratio_pct)}%, net credit demand {f(r.credit_demand_net_pct)} points (banks reporting higher demand minus lower).",
            f"- Rank range across five weighting scenarios: {int(r.rank_min)} to {int(r.rank_max)}.", "", f"**Evidence.** {why}", "", "**Conversation starters.**"] + [f"- {q}" for q in qs] + ["", f"**How to use it.** {so}", ""]
open(os.path.join(ROOT, "briefs", "industry_briefs.md"), "w").write("\n".join(out))

top = S.index[:3].tolist()
art = f"""# Where should Kenyan businesses expect growth in FY27? Ten industries, one scorecard

*Draft thought-leadership article (about 600 words). Partner review required before publishing. All figures from KNBS and CBK publications listed at the end.*

Kenya's economy grew 5.4% in the first quarter of 2026 (all industries, real terms), after 4.8% in 2025. Behind the average, sectors are moving at very different speeds, and banks are reading them differently. We combined the latest official numbers into one scorecard for ten industry groups.

**What we measured.** Five things: the size of each sector, its recent growth, its four-year growth record, how many credit officers see higher demand for credit, and how healthy bank lending to the sector is.

**What stood out.**

1. **{top[0]} leads.** Its output grew 6.5% in 2025 and 6.3% in Q1 2026, it has the cleanest loan book of any bank sector (3.3% non-performing, against 16.0% for the whole banking sector), and 38% of banks see higher demand for credit.
2. **{top[1]} is where demand for credit is strongest.** Nearly two-thirds of banks (65%) say demand from trade is up. But one in five trade loans (19.9%) is non-performing, so growth and risk travel together.
3. **{top[2]} is large and growing** (about 12% of the economy), though the credit signal is missing, so we rate it on size and growth alone.
4. **Real estate and construction carry the most stress.** Cement consumption rose 20.3% and construction output 9.4% in 2025, yet 24.1% of loans to real estate and 29.5% to building and construction are non-performing.
5. **Agriculture is big but slow.** It is 18.3% of GDP and grew 3.1% in 2025. Horticulture exports rose 13.8% while tea fell 8.0%.

**Does the ranking hold?** We re-ran it under five weighting schemes (growth-led, risk-aware, size-led, equal, balanced). Financial services stays first in all five, and consumer & retail stays between second and fourth. Agribusiness moves the most (second to seventh).

**What it means for leaders.** Growth is broad but uneven, and credit risk is concentrated in sectors that are also building. Businesses planning FY27 should test their plans against sector-specific risk, not the headline growth rate.

**Limits.** Our groupings are analytical, not official; two groups share one bank sector; two have no matching bank data; and the latest GDP figures cover Q1 2026 only.

*Sources: KNBS Quarterly GDP Report Q1 2026; KNBS Economic Survey 2026; CBK Bank Supervision Annual Report 2025; CBK Credit Officer Survey, June 2026.*
"""
open(os.path.join(ROOT, "comms", "thought_leadership_article.md"), "w").write(art)

posts = [("2026-10-26","Mon","LinkedIn","Thought leadership","All industries","Kenya grew 5.4% in Q1 2026. But the ten industry groups moved at very different speeds. Our scorecard ranks them. Article link in comments.","TL-01"),
 ("2026-10-27","Tue","X","Chart","Financial Services","Financial services: 6.3% growth in Q1 2026 and the lowest bad-loan ratio of any bank sector (3.3%). What does that mean for lenders?","IN-01"),
 ("2026-10-28","Wed","LinkedIn","Industry insight","Consumer & Retail","65% of banks see higher credit demand from trade, yet 19.9% of trade loans are non-performing. Growth and risk move together.","IN-02"),
 ("2026-10-29","Thu","LinkedIn","Event promotion","All industries","Save the date: Annual Industry Conference, 19 November. We will share the ranking and take your questions.","EV-01"),
 ("2026-10-30","Fri","X","Insight","Agribusiness","Agriculture is 18.3% of GDP but grew 3.1% in 2025. Horticulture exports +13.8%, tea -8.0%.","IN-03"),
 ("2026-11-02","Mon","LinkedIn","Industry insight","Real Estate & Construction","Cement consumption +20.3%, construction output +9.4% in 2025. Yet 29.5% of construction loans are non-performing.","IN-04"),
 ("2026-11-03","Tue","X","Chart","Technology, Media & Telecom","Mobile subscriptions reached 78.4 million in 2025 (+9.8%), but mobile money transfers fell 5.3%.","IN-05"),
 ("2026-11-04","Wed","LinkedIn","Event promotion","All industries","Speakers announced for the Annual Industry Conference.","EV-02"),
 ("2026-11-05","Thu","LinkedIn","Industry insight","Energy & Resources","Electricity consumption rose 9.6% to 11,785 GWh in 2025. Energy and water lending has a 6.8% NPL ratio.","IN-06"),
 ("2026-11-06","Fri","X","Insight","Public Sector & Development","Revenue grew 15.6% in 2025, spending 4.7%. External debt service takes 31.1% of export earnings.","IN-07"),
 ("2026-11-09","Mon","LinkedIn","Method note","All industries","How we built the scorecard: five pillars, five weighting scenarios, and what it cannot tell you.","TL-02"),
 ("2026-11-10","Tue","X","Insight","Manufacturing","Manufacturing grew 2.0% in 2025, the slowest of ten groups, while credit demand is firm (net +30).","IN-08"),
 ("2026-11-12","Thu","LinkedIn","Event promotion","All industries","One week to go: register for the Annual Industry Conference.","EV-03"),
 ("2026-11-16","Mon","LinkedIn","Event promotion","All industries","Agenda and what to expect.","EV-04"),
 ("2026-11-19","Thu","LinkedIn","Event live","All industries","Live from the Annual Industry Conference.","EV-05"),
 ("2026-11-20","Fri","LinkedIn","Event follow-up","All industries","Thank you. Three themes from the day.","EV-06")]
wb = Workbook(); ws = wb.active; ws.title = "Calendar"
hd = ["Date","Day","Channel","Pillar","Audience","Post copy (draft, figures from the scorecard)","Template ID","Approver","Status"]; ws.append(hd)
for c in ws[1]: c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="12202B"); c.alignment = Alignment(wrap_text=True)
for p in posts: ws.append([pd.Timestamp(p[0]).to_pydatetime()] + list(p[1:]) + ["Industry partner" if p[3] in ("Industry insight","Insight","Chart") else "Partner review", "Draft"])
for row in ws.iter_rows(min_row=2):
    row[0].number_format = "yyyy-mm-dd"; row[5].alignment = Alignment(wrap_text=True, vertical="top")
    for c in row: c.font = Font(name="Arial", size=10)
dv = DataValidation(type="list", formula1='"Idea,Draft,Partner review,Approved,Scheduled,Published"'); ws.add_data_validation(dv); dv.add(f"I2:I{len(posts)+1}")
for k, v in dict(A=12,B=6,C=10,D=18,E=26,F=90,G=11,H=18,I=14).items(): ws.column_dimensions[k].width = v
ws.freeze_panes = "A2"
s2 = wb.create_sheet("Status"); s2.append(["Status","Posts"])
for i, st in enumerate(["Idea","Draft","Partner review","Approved","Scheduled","Published"]): s2.append([st, f"=COUNTIF(Calendar!$I$2:$I${len(posts)+1},A{i+2})"])
s2.append(["Total", "=SUM(B2:B7)"])
wb.save(os.path.join(ROOT, "comms", "content_calendar.xlsx")); print("comms ok")
